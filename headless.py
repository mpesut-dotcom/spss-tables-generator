"""Headless/API ulaz za Hendalice: pokreni spremljeni _po.json plan obrade bez Streamlita.

CLI (isto sto i gumb "Generiraj tablice" u app.py, ukljucujuci skriveni _AI_META sheet
kad plan kaze global.add_ai_meta=true ili --ai-meta on):

    python headless.py --sav data_v5.sav --input input.txt --po data_v5_po.json
                       --output tables.xlsx [--design hendal] [--btw auto|on|off]
                       [--toc auto|on|off] [--ai-meta auto|on|off]

Ili iz Pythona::

    from headless import run
    errors = run(sav, input_txt, po_json, output)

Odnos prema ostatku repoa:
- engine funkcije dolaze iz spss_tables.py, orkestracijski helperi iz app.py
  (app.py je safe za import: sve je iza __main__ guarda, streamlit radi u bare modu);
- generacijska petlja zivi u app.py-jevom button handleru i ne moze se importati -
  ovaj modul je zrcali (izvor istine: app.py ~3790-4520). Svaka promjena te petlje
  u app.py mora se preslikati ovdje (i obratno).
- _AI_META redove pise ai_meta.AiMetaWriter kroz hookove (start / begin_output /
  add_total_sheet / add_between_options / add_krizanje_output / add_banner_table / finish)
  koje app.py i ovaj modul zovu na istim mjestima petlje - ovdje nema kopije nijednog reda.
  Study polja i routing dolaze iz plana (global.ai_meta, kako ga GUI snimi u _po.json).
- _po.json table_indices su POZICIONALNI u input.txt: ako se input.txt mijenja
  nakon snimanja plana, indeksi zastare. Out-of-range indeksi se toleriraju kao u
  app.py, uz WARNING.

Povijest: nastao kao po_runner.py u "agent hendal" delegacijskom sustavu
(validiran zero-diff protiv GUI outputa na Digitalni identiteti replayu);
upstreaman ovamo 2026-08-31 uz dozvolu vlasnika da GUI i automatika dijele kod;
2026-09-07 _AI_META port (ai_meta.py, zajednicki writer za oba ulaza).
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile

DEFAULT_ENGINE_DIR = os.path.dirname(os.path.abspath(__file__))


def _import_engine(engine_dir: str):
    if engine_dir not in sys.path:
        sys.path.insert(0, engine_dir)
    import spss_tables as eng  # noqa: E402
    import app as gui  # noqa: E402  (headless: defs only behind __main__ guard)
    import ai_meta as aim  # noqa: E402  (the _AI_META writer, shared with app.py)

    return eng, gui, aim


def _copy_sheet(src_ws, dest_ws) -> None:
    """Cell-by-cell copy incl. styles/widths/merges — app.py's TOTAL sheet transfer."""
    for row in src_ws.iter_rows():
        for cell in row:
            dest_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                d = dest_ws.cell(row=cell.row, column=cell.column)
                d.font = cell.font.copy()
                d.fill = cell.fill.copy()
                d.border = cell.border.copy()
                d.alignment = cell.alignment.copy()
                d.number_format = cell.number_format
    for col_letter, dim in src_ws.column_dimensions.items():
        dest_ws.column_dimensions[col_letter].width = dim.width
    for mc in src_ws.merged_cells.ranges:
        dest_ws.merge_cells(str(mc))


def _build_toc_rows(titles, eng):
    """Question-centric TOC structure — app.py ~3839-3890."""
    toc_rows: list[dict] = []
    base_to_row: dict[str, int] = {}
    mean_pending: list[tuple[int, str]] = []

    for ti in range(len(titles)):
        ttype = eng.get_table_type(titles[ti])
        ttitle = eng.get_table_title(titles[ti])
        stripped = ttitle.rstrip()
        if ttype in ("n", "m"):
            mv = re.match(r"^(\S+)", ttitle.strip())
            mean_pending.append((ti, mv.group(1).rstrip(".:") if mv else ""))
        elif "- T2B" in stripped:
            base = stripped[: stripped.index("- T2B")].rstrip()
            if base in base_to_row:
                toc_rows[base_to_row[base]]["t2b_idx"] = ti
            else:
                base_to_row[base] = len(toc_rows)
                toc_rows.append({"title": base, "regular_idx": None, "mean_idx": None, "t2b_idx": ti})
        else:
            if stripped not in base_to_row:
                base_to_row[stripped] = len(toc_rows)
                toc_rows.append({"title": stripped, "regular_idx": ti, "mean_idx": None, "t2b_idx": None})

    for mi, mvar in mean_pending:
        matched = False
        if mvar:
            for row in toc_rows:
                rm = re.match(r"^(\S+)", row["title"].strip())
                rv = rm.group(1).rstrip(".:") if rm else ""
                if rv and (rv == mvar or (rv.startswith(mvar) and len(rv) > len(mvar) and rv[len(mvar)] == ".")):
                    row["mean_idx"] = mi
                    matched = True
        if not matched:
            mt = eng.get_table_title(titles[mi]).rstrip()
            if "- MEAN" in mt:
                mt = mt[: mt.index("- MEAN")].rstrip()
            if mt not in base_to_row:
                base_to_row[mt] = len(toc_rows)
                toc_rows.append({"title": mt, "regular_idx": None, "mean_idx": mi, "t2b_idx": None})
    return toc_rows


def _write_toc(wb, toc_rows, toc_positions, toc_kriz_sheets, titles, eng, design):
    """TOC (BETA) sheet — app.py ~4376-4506."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    theme = eng._get_theme(design)
    ws = wb.create_sheet("TOC (BETA)", 0)

    hfont = Font(name="Calibri", size=11, bold=True, color=theme.get("toc_header_color", theme["title_color"]))
    hfill = PatternFill(start_color=theme["toc_header_fill"], end_color=theme["toc_header_fill"], fill_type="solid")
    lfont = Font(name="Calibri", size=10, color=theme["toc_link_color"], underline="single")
    dfont = Font(name="Calibri", size=10, color=theme["data_color"])
    even_fill = PatternFill(start_color=theme["toc_even_fill"], end_color=theme["toc_even_fill"], fill_type="solid")
    brd = Border(bottom=Side(style="thin", color=theme["toc_line_color"]))
    hbrd = Border(bottom=Side(style="medium", color="1D1D1B"))

    def var_name(title):
        m = re.match(r"^(\S+)", title.strip())
        return m.group(1).rstrip(".:") if m else title[:20]

    def make_link(row, col, text, sheet, cell_row, fill=None):
        safe = sheet.replace("'", "''")
        c = ws.cell(row=row, column=col, value=text)
        c.hyperlink = f"#'{safe}'!A{cell_row}"
        c.font = lfont
        c.border = brd
        if fill:
            c.fill = fill
        c.alignment = Alignment(horizontal="center")

    def blank(row, col, fill):
        c = ws.cell(row=row, column=col)
        c.border = brd
        if fill:
            c.fill = fill

    kriz_cols = []
    for ki, ks in enumerate(toc_kriz_sheets):
        kriz_cols.append((ks["plain"], ki, "kriz"))
        if ks["sig"]:
            kriz_cols.append((ks["sig"], ki, "kriz_sig"))
        if ks["sigT"]:
            kriz_cols.append((ks["sigT"], ki, "kriz_sigT"))
    all_headers = ["#", "Pitanje", "Total", "MEAN", "T2B"] + [kc[0] for kc in kriz_cols]

    for ci, h in enumerate(all_headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = hbrd

    for ri, qrow in enumerate(toc_rows, 2):
        row_fill = even_fill if ri % 2 == 0 else None
        c = ws.cell(row=ri, column=1, value=ri - 1)
        c.font = dfont
        c.border = brd
        if row_fill:
            c.fill = row_fill
        c.alignment = Alignment(horizontal="center")

        c = ws.cell(row=ri, column=2, value=qrow["title"])
        c.font = dfont
        c.border = brd
        if row_fill:
            c.fill = row_fill

        for col, idx_key, link_title in (
            (3, "regular_idx", qrow["title"]),
            (4, "mean_idx", None),
            (5, "t2b_idx", qrow["title"]),
        ):
            idx = qrow.get(idx_key)
            tp = [p for p in toc_positions.get(idx, []) if p["cat"] == "total"] if idx is not None else []
            if tp:
                text = var_name(link_title if link_title is not None else eng.get_table_title(titles[idx]))
                make_link(ri, col, text, tp[0]["sheet"], tp[0]["row"], row_fill)
            else:
                blank(ri, col, row_fill)

        for kci, (kname, kidx, kcat) in enumerate(kriz_cols):
            col_num = 6 + kci
            found = False
            for try_idx in [qrow.get("regular_idx"), qrow.get("t2b_idx"), qrow.get("mean_idx")]:
                if try_idx is None:
                    continue
                tp = [p for p in toc_positions.get(try_idx, []) if p["cat"] == kcat and p.get("kriz_idx") == kidx]
                if tp:
                    make_link(ri, col_num, var_name(qrow["title"]), tp[0]["sheet"], tp[0]["row"], row_fill)
                    found = True
                    break
            if not found:
                blank(ri, col_num, row_fill)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 55
    for ci in range(3, len(all_headers) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 14
    ws.freeze_panes = "C2"


def run(
    sav: str,
    input_txt: str,
    po_json: str,
    output: str,
    engine_dir: str = DEFAULT_ENGINE_DIR,
    design: str = "hendal",
    btw: str = "auto",
    toc: bool | None = None,
    ai_meta: bool | None = None,
) -> list[str]:
    """Replay *po_json* on *sav* + *input_txt* into *output*. Returns warnings/errors."""
    import pyreadstat
    from openpyxl import Workbook, load_workbook

    eng, gui, aim = _import_engine(engine_dir)

    df, meta = pyreadstat.read_sav(sav, apply_value_formats=False)
    break_vars, titles, variables = eng.parse_input_file(input_txt)
    with open(po_json, encoding="utf-8") as fh:
        po = json.load(fh)

    g = po.get("global", {})
    use_weight = g.get("use_weight", False)
    weight_col = g.get("weight_col") if use_weight else None
    start_num = g.get("start_num", 1)
    add_toc = g.get("add_toc", True) if toc is None else toc
    add_ai_meta = g.get("add_ai_meta", False) if ai_meta is None else ai_meta
    global_fgs = [fg for fg in g.get("filter_groups", []) if fg.get("vals")]
    output_defs = po.get("outputs", [])

    all_errors: list[str] = []
    n_tables = len(titles)
    for od in output_defs:
        stale = [i for i in od.get("table_indices", []) if i >= n_tables]
        if stale:
            all_errors.append(
                f"WARNING output '{od.get('sheet_name')}': {len(stale)} stale table_indices "
                f">= {n_tables} tables in input.txt (po.json older than input.txt?) — skipped like app.py"
            )

    wb = Workbook()
    wb.remove(wb.active)
    existing_sheets: list[str] = [aim.AI_META_SHEET_NAME] if add_ai_meta else []

    toc_rows = _build_toc_rows(titles, eng) if add_toc else []
    toc_positions: dict[int, list[dict]] = {}
    toc_kriz_sheets: list[dict] = []

    meta_writer = None
    if add_ai_meta:
        meta_writer = aim.AiMetaWriter.start(
            df, meta, titles, variables,
            plan=g.get("ai_meta") or {},
            use_weight=use_weight, weight_col=weight_col, start_num=start_num, table_design=design,
            sav_name=os.path.basename(sav), input_name=os.path.basename(input_txt),
            global_filter_groups=global_fgs,
        )

    def unique_name(base: str) -> str:
        name = base[:31]
        while name in existing_sheets:
            sfx = 2
            while f"{name[:28]}_{sfx}" in existing_sheets:
                sfx += 1
            name = f"{name[:28]}_{sfx}"
        existing_sheets.append(name)
        return name

    for out_i, out_def in enumerate(output_defs):
        output_id = out_i + 1
        work_df = df.copy()
        if global_fgs:
            work_df = gui.apply_filter_groups(work_df, global_fgs)
        n_after_global_filter = len(work_df)
        if out_def.get("filter_groups"):
            work_df = gui.apply_filter_groups(work_df, out_def["filter_groups"])
        if meta_writer is not None:
            meta_writer.begin_output(output_id, out_def, n_after_global_filter, len(work_df))

        if out_def["type"] == "total":
            tbl_indices_set = set(out_def.get("table_indices", []))
            tables, errs = gui.generate_tables(work_df, meta, titles, variables, weight_col, start_num)
            if tbl_indices_set:
                tables = [t for t in tables if t.get("_idx") in tbl_indices_set]
            all_errors.extend(errs)

            show_btw = {"on": True, "off": False}.get(btw, out_def.get("show_between_options", True))
            between_blocks = []
            if show_btw:
                bo_col_map = eng.build_column_map(work_df)
                for btbl in tables:
                    bidx = btbl.get("_idx")
                    if bidx is None:
                        continue
                    bo = eng.compute_between_options(
                        work_df, eng.get_table_type(titles[bidx]), variables[bidx], btbl, meta, bo_col_map, weight_col
                    )
                    if bo:
                        between_blocks.append(
                            {
                                "title": eng.get_table_title(titles[bidx]),
                                "q_code": gui._extract_q_code_from_title(eng.get_table_title(titles[bidx])),
                                "table_type": eng.get_table_type(titles[bidx]),
                                "table_idx": bidx,
                                "results": bo,
                            }
                        )

            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
                tf_path = tf.name
            eng.write_tables_to_excel(tables, tf_path, design=design)
            twb = load_workbook(tf_path)
            for src_ws in twb.worksheets:
                sname = unique_name(out_def["sheet_name"])
                dest_ws = wb.create_sheet(title=sname)
                _copy_sheet(src_ws, dest_ws)
                if meta_writer is not None:
                    meta_writer.add_total_sheet(output_id, sname, out_def, tables)
                if add_toc:
                    toc_r = 1
                    for tbl in tables:
                        gi = tbl.get("_idx")
                        if gi is not None:
                            toc_positions.setdefault(gi, []).append({"sheet": sname, "row": toc_r, "cat": "total"})
                        toc_r += 1 + 1 + len(tbl["rows"]) + (1 if tbl.get("caption", "") else 0) + 1
            twb.close()
            os.unlink(tf_path)

            if between_blocks:
                btw_name = unique_name(out_def["sheet_name"][:27] + "_btw")
                btw_ws = eng.write_between_options_sheet(wb, btw_name, between_blocks, design=design)
                if btw_ws is not None and meta_writer is not None:
                    meta_writer.add_between_options(output_id, sname, btw_name, between_blocks)

        elif out_def["type"] == "krizanje":
            banner_vars = out_def.get("banner_vars", [])
            tbl_indices = out_def.get("table_indices", [])
            show_sig = out_def.get("show_sig", True)
            show_sig_total = out_def.get("show_sig_total", False)
            if not banner_vars or not tbl_indices:
                all_errors.append(f"Output '{out_def['sheet_name']}': nema banner varijabli ili tablica")
                continue

            col_map = eng.build_column_map(work_df)
            banner_labels_by_var = {bv: eng.get_var_label(bv, meta) for bv in banner_vars}

            banner_entries = []
            for ti in tbl_indices:
                entry, entry_errors = gui._build_banner_table_entry(
                    work_df, meta, col_map, titles, variables, ti, banner_vars, banner_labels_by_var, weight_col, start_num
                )
                all_errors.extend(entry_errors)
                if entry is not None:
                    banner_entries.append(entry)
            if not banner_entries:
                continue

            ws = wb.create_sheet(title=unique_name(out_def["sheet_name"]))
            ws_sig = wb.create_sheet(title=unique_name(out_def["sheet_name"][:27] + "_sig")) if show_sig else None
            ws_sig_total = (
                wb.create_sheet(title=unique_name(out_def["sheet_name"][:21] + "_sig_total")) if show_sig_total else None
            )

            if meta_writer is not None:
                meta_writer.add_krizanje_output(
                    output_id, out_def, ws.title,
                    ws_sig.title if ws_sig is not None else "",
                    ws_sig_total.title if ws_sig_total is not None else "",
                    len(banner_entries), banner_vars, [banner_labels_by_var[bv] for bv in banner_vars],
                    show_sig, show_sig_total, work_df,
                )

            if add_toc:
                ki = len(toc_kriz_sheets)
                toc_kriz_sheets.append(
                    {
                        "plain": ws.title,
                        "sig": ws_sig.title if ws_sig is not None else None,
                        "sigT": ws_sig_total.title if ws_sig_total is not None else None,
                    }
                )

            current_row = current_row_sig = current_row_st = 1
            for entry in banner_entries:
                ti = entry["table_idx"]
                if add_toc:
                    toc_positions.setdefault(ti, []).append(
                        {"sheet": ws.title, "row": current_row, "cat": "kriz", "kriz_idx": ki}
                    )
                    if ws_sig is not None:
                        toc_positions.setdefault(ti, []).append(
                            {"sheet": ws_sig.title, "row": current_row_sig, "cat": "kriz_sig", "kriz_idx": ki}
                        )
                    if ws_sig_total is not None:
                        toc_positions.setdefault(ti, []).append(
                            {"sheet": ws_sig_total.title, "row": current_row_st, "cat": "kriz_sigT", "kriz_idx": ki}
                        )
                plain_start = current_row
                current_row = eng.write_banner_to_sheet(
                    ws, entry["banner"], entry["title"], start_row=current_row, show_sig=False,
                    design=design, banner_labels=entry["banner_labels"],
                )
                if meta_writer is not None:
                    meta_writer.add_banner_table(output_id, ws.title, "cross_base", entry, plain_start, current_row - 1)
                current_row += 2
                if ws_sig is not None:
                    sig_start = current_row_sig
                    current_row_sig = eng.write_banner_to_sheet(
                        ws_sig, entry["banner"], entry["title"], start_row=current_row_sig, show_sig=True,
                        design=design, banner_labels=entry["banner_labels"],
                    )
                    if meta_writer is not None:
                        meta_writer.add_banner_table(output_id, ws_sig.title, "significance", entry, sig_start, current_row_sig - 1)
                    current_row_sig += 2
                if ws_sig_total is not None:
                    sigt_start = current_row_st
                    current_row_st = eng.write_banner_to_sheet(
                        ws_sig_total, entry["banner"], entry["title"], start_row=current_row_st, show_sig=True,
                        show_sig_total=True, design=design, banner_labels=entry["banner_labels"],
                    )
                    if meta_writer is not None:
                        meta_writer.add_banner_table(output_id, ws_sig_total.title, "sig_total", entry, sigt_start, current_row_st - 1)
                    current_row_st += 2
        else:
            all_errors.append(f"Output '{out_def.get('sheet_name')}': nepoznat tip '{out_def.get('type')}'")

    if add_toc and toc_rows:
        _write_toc(wb, toc_rows, toc_positions, toc_kriz_sheets, titles, eng, design)

    if meta_writer is not None:
        meta_writer.finish(wb)

    if not wb.worksheets or not any(ws.sheet_state == "visible" for ws in wb.worksheets):
        wb.create_sheet("Sheet1")
    wb.save(output)
    wb.close()
    return all_errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--sav", required=True)
    parser.add_argument("--input", required=True)
    parser.add_argument("--po", required=True, help="_po.json spremljen iz Hendalice GUI-ja (ili rucno napisan)")
    parser.add_argument("--output", required=True)
    parser.add_argument("--engine-dir", default=DEFAULT_ENGINE_DIR)
    parser.add_argument("--design", default="hendal")
    parser.add_argument("--btw", choices=["auto", "on", "off"], default="auto",
                        help="between-options sig sheets (auto = po po.json/app defaultu)")
    parser.add_argument("--toc", choices=["auto", "on", "off"], default="auto")
    parser.add_argument("--ai-meta", choices=["auto", "on", "off"], default="auto",
                        help="skriveni _AI_META sheet za AI context exporter (auto = global.add_ai_meta iz po.json)")
    args = parser.parse_args()

    toc = None if args.toc == "auto" else args.toc == "on"
    ai_meta = None if args.ai_meta == "auto" else args.ai_meta == "on"
    errors = run(args.sav, args.input, args.po, args.output, args.engine_dir, args.design, args.btw, toc, ai_meta)
    for err in errors:
        print(f"  ! {err}")
    print(f"Gotovo: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
