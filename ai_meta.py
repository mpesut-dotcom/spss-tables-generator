#!/usr/bin/env python3
"""_AI_META — the metadata sheet the AI context exporter reads (visible since 2026-10-07, the owner's word: the researcher sees what the report gate reads; hidden before) (Raspisivanje slajdova, schema v3.4).

ONE writer for both entry points. app.py ("Generiraj tablice") and headless.py (po.json replay)
call the same AiMetaWriter hooks at the same points of their generate loops; neither file holds
a copy of a META row. Change a row here and both outputs change; add a hook here and call it
from both loops (headless.py's docstring carries the mirror rule).

The sheet (AiMetaBuilder.write_to_workbook): one block per section — a `## SECTION` line, a
header row, the rows, a blank line. META, SHEETS and TABLES are always written, the other
sections only when they have rows. Section order: META, STUDY_CONTEXT, SHEETS, OUTPUTS, FILTERS,
TABLES, OUTPUT_TABLES, BANNERS, BETWEEN_OPTIONS, ROUTING, WARNINGS. The exporter requires META,
SHEETS, OUTPUTS, TABLES, OUTPUT_TABLES, BANNERS and meta_schema_version major 1.

Contents, top to bottom:
- shared parsing/description helpers app.py also uses — moved here so this module never
  imports app.py: _extract_vars_from_line, build_filter_groups_description, _NumpyEncoder,
  _json_cell, _excel_cell_value, _normalize_user_text
- AiMetaBuilder (the sheet writer) and the per-section row helpers, verbatim from app.py
  (ported 2026-09-07)
- AiMetaWriter — the hooks, at the bottom

The Streamlit form and its plan collection (_collect_ai_meta_plan / _apply_ai_meta_plan) stay
in app.py; the plan they produce is the `global.ai_meta` dict of the saved _po.json, which
headless.py hands to AiMetaWriter.start unchanged — so GUI and headless write the same rows
for the same plan (only `generated_at` differs).
"""

import json
import os
import re
from datetime import datetime, timezone

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

from spss_tables import get_table_title, get_table_type, get_var_label

def _extract_vars_from_line(var_line):
    """Izvuci pojedinačne varijable iz var definicijskog reda."""
    def _unique_tokens(tokens):
        result = []
        seen = set()
        for token in tokens:
            key = token.lower()
            if key not in seen:
                seen.add(key)
                result.append(token)
        return result

    var_line = var_line.strip()
    # MR: $e1 '' var1 var2 var3
    if var_line.startswith('$'):
        return _unique_tokens([p for p in var_line.split() if not p.startswith('$') and p != "''"])
    # Numeric composite: q1_1 q1_2 ... q1_1+q1_2+...
    if '+' in var_line:
        return _unique_tokens([p for p in var_line.split() if '+' not in p])
    # Simple or single-var MEAN line, e.g. "q2a" or "q2a q2a".
    return _unique_tokens(var_line.split()) if var_line else []


def build_filter_groups_description(filter_groups, labels_dict, val_labels_dict):
    """Napravi citljiv opis filter grupa."""
    def _short_text(value, limit):
        text = '' if value is None else str(value)
        return text[:limit]

    def _lookup_value_label(vlabels, value):
        if not hasattr(vlabels, 'get'):
            return ''

        candidates = [value]
        try:
            float_value = float(value)
            candidates.append(float_value)
            if float_value.is_integer():
                candidates.append(int(float_value))
        except (ValueError, TypeError):
            pass

        candidates.append(str(value))

        for candidate in candidates:
            try:
                lbl = vlabels.get(candidate, '')
            except TypeError:
                continue
            if lbl:
                return lbl
        return ''

    def _lookup_multi_value_label(group_vars, value):
        for group_var in group_vars:
            lbl = _lookup_value_label(val_labels_dict.get(group_var, {}), value)
            if lbl:
                return lbl
        return value

    parts = []
    for i, grp in enumerate(filter_groups):
        vals = grp.get('vals', [])
        if not vals:
            continue

        mode = grp.get('mode', 'single')
        logic = grp.get('logic', 'AND')
        group_label = grp.get('group_label', '')

        if mode == 'multi':
            group_vars = grp.get('vars', [])
            val_parts = [
                _short_text(_lookup_multi_value_label(group_vars, value), 40)
                for value in vals
            ]
            val_str = ' ILI '.join(val_parts)
            part = f"`{_short_text(group_label, 35)}` = {val_str}"
        else:
            var = grp['var']
            var_lbl = labels_dict.get(var) or var
            short_var = _short_text(var_lbl, 35) if var_lbl != var else str(var)

            vlabels = val_labels_dict.get(var, {})
            val_parts = []
            for v in vals:
                lbl = vlabels.get(v, '')
                if not lbl:
                    try:
                        lbl = vlabels.get(float(v), '')
                    except (ValueError, TypeError):
                        pass
                if not lbl:
                    try:
                        lbl = vlabels.get(int(float(v)), '')
                    except (ValueError, TypeError):
                        pass
                val_parts.append(str(lbl) if lbl else str(v))

            val_str = ' ILI '.join(val_parts)
            part = f"`{short_var}` = {val_str}"

        if i > 0 and parts:
            connector = ' **ILI** ' if logic == 'OR' else ' **I** '
            parts.append(connector)
        parts.append(part)

    return ''.join(parts) if parts else ''


class _NumpyEncoder(json.JSONEncoder):
    """Handle numpy types when serializing to JSON."""
    def default(self, obj):
        if isinstance(obj, (np.integer,)):
            return int(obj)
        if isinstance(obj, (np.floating,)):
            return float(obj)
        if isinstance(obj, np.ndarray):
            return obj.tolist()
        return super().default(obj)


def _json_cell(value):
    """Serialize structured values for storage in a worksheet cell."""
    if value is None:
        return ''
    return json.dumps(value, ensure_ascii=False, cls=_NumpyEncoder)


def _excel_cell_value(value):
    """Convert numpy/pandas-ish values into values openpyxl can store."""
    if value is None:
        return ''
    try:
        if pd.isna(value):
            return ''
    except (TypeError, ValueError):
        pass
    if isinstance(value, (dict, list, tuple, set)):
        return _json_cell(list(value) if isinstance(value, set) else value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, np.ndarray):
        return _json_cell(value.tolist())
    if isinstance(value, bool):
        return value
    return value


def _normalize_user_text(value):
    if value is None:
        return ''
    return str(value).replace('\r\n', '\n').replace('\r', '\n').replace('\\n', '\n')


class AiMetaBuilder:
    """Collect and write parser-friendly metadata for AI context export."""
    SECTION_COLUMNS = {
        'META': ['key', 'value', 'source', 'notes'],
        'STUDY_CONTEXT': ['key', 'value', 'source', 'client_facing', 'notes'],
        'SHEETS': ['sheet_name', 'role', 'output_id', 'base_sheet_name',
                   'table_block_count', 'hidden'],
        'OUTPUTS': ['output_id', 'sheet_base_name', 'output_type',
                    'actual_sheet_base', 'actual_sheet_sig',
                    'actual_sheet_sig_total', 'selected_table_indices_json',
                    'selected_table_count', 'banner_vars_json',
                    'banner_labels_json', 'show_sig', 'show_sig_total',
                    'global_filter_description', 'output_filter_description',
                    'effective_filter_description', 'filter_groups_json',
                    'global_filter_groups_json', 'n_unfiltered',
                    'n_after_global_filter', 'n_after_effective_filter',
                    'weight_col'],
        'FILTERS': ['scope', 'output_id', 'condition_order',
                    'logic_with_previous', 'mode', 'var', 'vars_json',
                    'group_label', 'var_label', 'value_codes_json',
                    'value_labels_json', 'description'],
        'TABLES': ['table_idx', 'table_number', 'q_code', 'title',
                   'input_title_line', 'input_var_line', 'table_type_code',
                   'table_type_label', 'variables_json', 'variable_labels_json',
                   'metric_type', 'result_level', 'answer_scale_json',
                   'evidence_needed_rules_json', 'valid_comparisons_json',
                   'invalid_comparisons_json', 'interpretation_limit_hints_json',
                   'is_multi_response', 'is_mean', 'is_t2b', 'full_base_n',
                   'full_base_status', 'routing_filter_description',
                   'routing_filter_expression', 'routing_filter_source',
                   'routing_filter_status'],
        'OUTPUT_TABLES': ['output_id', 'sheet_name', 'sheet_role', 'table_idx',
                          'table_number', 'title_rendered', 'start_row',
                          'end_row', 'base_n', 'base_source',
                          'effective_filter_description',
                          'effective_filter_groups_json',
                          'routing_filter_description',
                          'routing_filter_expression', 'routing_filter_status',
                          'base_frame',
                          'banner_vars_json'],
        'BANNERS': ['output_id', 'sheet_name', 'banner_order', 'banner_var',
                    'banner_label', 'axis_id', 'segment_order', 'segment_code',
                    'segment_label', 'segment_n', 'weight_col'],
        'BETWEEN_OPTIONS': ['output_id', 'source_sheet', 'between_sheet',
                            'table_idx', 'table_number', 'q_code', 'title',
                            'table_type_code', 'metric_type', 'scope', 'axis',
                            'base_n', 'option_a_label', 'option_a_value',
                            'option_b_label', 'option_b_value', 'significant',
                            'direction', 'test', 'confidence', 'b_a_not_b',
                            'c_b_not_a', 'note'],
        'ROUTING': ['table_idx', 'table_number', 'q_code', 'title', 'base_n',
                    'description', 'expression', 'source'],
        'WARNINGS': ['level', 'code', 'table_idx', 'output_id', 'message',
                     'audience', 'client_exclude'],
    }

    def __init__(self):
        self.sections = {section: [] for section in self.SECTION_COLUMNS}

    def add(self, section, **row):
        if section not in self.sections:
            raise KeyError(section)
        self.sections[section].append(row)

    def write_to_workbook(self, wb, sheet_name='_AI_META', hidden=False):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        row_num = 1

        for section, columns in self.SECTION_COLUMNS.items():
            rows = self.sections.get(section, [])
            if not rows and section not in ('META', 'SHEETS', 'TABLES'):
                continue

            ws.cell(row=row_num, column=1, value=f'## {section}')
            row_num += 1
            for col_num, column_name in enumerate(columns, 1):
                ws.cell(row=row_num, column=col_num, value=column_name)
            row_num += 1

            for section_row in rows:
                for col_num, column_name in enumerate(columns, 1):
                    ws.cell(
                        row=row_num,
                        column=col_num,
                        value=_excel_cell_value(section_row.get(column_name, '')),
                    )
                row_num += 1
            row_num += 1

        for col_idx, width in enumerate((26, 28, 18, 45, 28, 28, 28, 18, 36, 36), 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = 'A3'
        if hidden:
            ws.sheet_state = 'hidden'
        return ws


_TABLE_TYPE_LABELS = {
    's': 'distribution',
    'k': 'multi_response',
    'd': 'multi_dichotomy',
    'n': 'numeric_full',
    'm': 'mean',
    'f': 'frequency',
}


_AI_META_STUDY_FIELDS = (
    'project_name', 'client', 'product', 'study_type', 'study_objective',
    'research_topics', 'fielding_period', 'default_time_scope',
    'default_geography', 'population', 'sample_design', 'sample_notes',
    'sample_size_notes', 'data_collection_method', 'methodology_notes',
    'weighting_notes', 'reporting_notes', 'methodology',
)


_AI_META_LEGACY_STUDY_FIELDS = (
    'project_name', 'study_type', 'methodology', 'fielding_period', 'client', 'product'
)


_AI_META_STUDY_DEFAULTS = {
    'default_geography': 'Hrvatska',
}


_AI_META_STUDY_NOTES = {
    'population': 'WHO: target population/universe for client-facing methodology.',
    'sample_notes': 'HOW: how the sample was constructed or controlled.',
    'methodology_notes': 'WHAT: data collection and method details, separate from population/sample.',
    'default_geography': 'Default geography used when question context has no narrower geography.',
    'default_time_scope': 'Default time scope used when question context has no narrower wave/period.',
    'weighting_notes': 'Client-readable weighting statement; never debug syntax.',
}


def _clean_inline_text(value):
    text = _normalize_user_text(value).strip()
    return re.sub(r'[ \t]+', ' ', text)


def _clean_period_text(value):
    text = _clean_inline_text(value)
    text = re.sub(
        r'\b([A-Za-zČĆŽŠĐčćžšđ]+),\s+(\d{4})(\.)?',
        lambda match: f"{match.group(1)} {match.group(2)}{match.group(3) or ''}",
        text,
    )
    return text


def _clean_product_text(value):
    text = _clean_inline_text(value)
    return re.sub(r'^Op[cć]enita tema\s*[-:]\s*', '', text, flags=re.IGNORECASE)


def _clean_study_type_text(value):
    text = _clean_inline_text(value)
    if re.search(r'\b(CAWI|CAPI|CATI|PAPI|computer assisted)\b', text, re.IGNORECASE) and ' - ' in text:
        return text.split(' - ', 1)[0].strip()
    return text


def _detect_collection_method(*texts):
    joined = ' '.join(_normalize_user_text(text) for text in texts if text)
    modes = []
    mode_labels = (
        ('CAWI', 'CAWI online anketa'),
        ('CAPI', 'CAPI osobni intervjui'),
        ('CATI', 'CATI telefonsko anketiranje'),
        ('PAPI', 'PAPI papirnata anketa'),
    )
    for code, label in mode_labels:
        if re.search(rf'\b{code}\b', joined, re.IGNORECASE):
            modes.append(label)
    return '; '.join(modes)


def _resolve_weighting_notes(study_meta, use_weight, weight_col):
    user_notes = _normalize_user_text(study_meta.get('weighting_notes', '')).strip()
    if user_notes:
        return user_notes, 'user'
    if use_weight and weight_col:
        return f'Podaci su ponderirani varijablom {weight_col}.', 'app'
    return 'Podaci nisu ponderirani.', 'app'


def _resolve_study_context(study_meta, df=None, use_weight=False, weight_col=None):
    raw = {field: _normalize_user_text(study_meta.get(field, '')) for field in _AI_META_STUDY_FIELDS}
    fielding_period = _clean_period_text(raw.get('fielding_period'))
    default_time_scope = _clean_period_text(raw.get('default_time_scope')) or fielding_period
    methodology_notes = _normalize_user_text(raw.get('methodology_notes')).strip() or _normalize_user_text(raw.get('methodology')).strip()
    data_collection_method = _clean_inline_text(raw.get('data_collection_method')) or _detect_collection_method(
        raw.get('study_type'), raw.get('methodology'), methodology_notes)
    weighting_notes, weighting_source = _resolve_weighting_notes(study_meta, use_weight, weight_col)

    context = {
        'project_name': _clean_inline_text(raw.get('project_name')),
        'client': _clean_inline_text(raw.get('client')),
        'product': _clean_product_text(raw.get('product')),
        'study_type': _clean_study_type_text(raw.get('study_type')),
        'study_objective': _normalize_user_text(raw.get('study_objective')).strip(),
        'research_topics': _clean_inline_text(raw.get('research_topics')),
        'fielding_period': fielding_period,
        'default_time_scope': default_time_scope,
        'default_geography': _clean_inline_text(raw.get('default_geography')) or _AI_META_STUDY_DEFAULTS['default_geography'],
        'population': _normalize_user_text(raw.get('population')).strip(),
        'sample_design': _normalize_user_text(raw.get('sample_design')).strip(),
        'sample_notes': _normalize_user_text(raw.get('sample_notes')).strip(),
        'sample_size_notes': _normalize_user_text(raw.get('sample_size_notes')).strip(),
        'data_collection_method': data_collection_method,
        'methodology_notes': methodology_notes,
        'weighting_notes': weighting_notes,
        'reporting_notes': _normalize_user_text(raw.get('reporting_notes')).strip(),
        'sample_n': int(len(df)) if df is not None else '',
        'weight_enabled': bool(use_weight),
        'weight_col': weight_col or '',
    }
    if use_weight and weight_col and df is not None and weight_col in df.columns:
        context['weighted_sample_n'] = round(float(df[weight_col].sum()), 1)
    else:
        context['weighted_sample_n'] = ''

    context['methodology'] = _normalize_user_text(raw.get('methodology')).strip() or context['methodology_notes']

    sources = {}
    for key, value in context.items():
        if key == 'weighting_notes':
            sources[key] = weighting_source
        elif key in ('sample_n', 'weight_enabled', 'weight_col', 'weighted_sample_n'):
            sources[key] = 'app'
        elif key == 'default_geography' and not _normalize_user_text(raw.get('default_geography')).strip():
            sources[key] = 'app_default'
        elif key == 'default_time_scope' and not _normalize_user_text(raw.get('default_time_scope')).strip() and fielding_period:
            sources[key] = 'fielding_period'
        elif key == 'data_collection_method' and not _normalize_user_text(raw.get('data_collection_method')).strip() and value:
            sources[key] = 'inferred'
        elif key == 'methodology_notes' and not _normalize_user_text(raw.get('methodology_notes')).strip() and value:
            sources[key] = 'legacy_methodology'
        else:
            sources[key] = 'user' if value not in ('', None) else 'unknown'
    return context, sources


def _add_meta_kv(builder, key, value, source='app', notes=''):
    builder.add('META', key=key, value=value, source=source, notes=notes)


def _add_study_meta_rows(builder, study_meta, df, use_weight, weight_col):
    study_context, sources = _resolve_study_context(study_meta, df, use_weight, weight_col)
    for key, value in study_context.items():
        source = sources.get(key, 'unknown')
        notes = _AI_META_STUDY_NOTES.get(key, '')
        if key in _AI_META_LEGACY_STUDY_FIELDS:
            _add_meta_kv(builder, key, value, source=source, notes=notes)
        _add_meta_kv(builder, f'study_context.{key}', value, source=source, notes=notes)
        builder.add(
            'STUDY_CONTEXT',
            key=key,
            value=value,
            source=source,
            client_facing=key not in ('weight_enabled', 'weight_col'),
            notes=notes,
        )
    _add_meta_kv(
        builder,
        'study_context_json',
        _json_cell(study_context),
        source='app',
        notes='Structured study context for downstream JSON exporter.',
    )


def _table_metric_type(table_type, table_title):
    title_upper = (table_title or '').upper()
    if table_type in ('n', 'm'):
        return 'mean_score'
    if table_type in ('k', 'd'):
        return 'multi_response'
    if 'T2B' in title_upper:
        return 't2b'
    if table_type == 'f':
        return 'frequency'
    return 'distribution'


def _table_result_level(full_base_status):
    if full_base_status == 'partial_base':
        return 'filtered_sample'
    if full_base_status == 'full_sample':
        return 'full_sample'
    if full_base_status == 'not_applicable':
        return 'not_applicable'
    return 'unknown'


def _answer_scale_meta(metric_type, variables_resolved, val_labels_dict):
    if metric_type == 'mean_score':
        return {'kind': 'numeric', 'source': 'numeric_variables'}
    values_by_label = []
    seen = set()
    for var_name in variables_resolved:
        value_labels = val_labels_dict.get(var_name, {}) or {}
        for code, label in sorted(
            value_labels.items(),
            key=lambda item: (0, float(item[0])) if isinstance(item[0], (int, float, np.integer, np.floating)) else (1, str(item[0])),
        ):
            label_text = _lookup_value_label(value_labels, code) or str(label)
            key = (str(code), label_text)
            if key not in seen:
                seen.add(key)
                values_by_label.append({'code': code, 'label': label_text})
    return {'kind': 'categorical', 'values': values_by_label}


def _table_evidence_needed_rules(metric_type):
    rules = {
        'total_answer': 'Use the total result and cite the available base N.',
        'segment_answer': 'Use crosstab segments only when the requested banner exists and base is sufficient.',
        'trend_answer': 'Use only comparable waves/time outputs that are present for this question.',
        'significance_answer': 'Use significance sheets when present; otherwise describe differences as directional, not statistically significant.',
    }
    if metric_type == 't2b':
        rules['total_answer'] = 'Use the aggregated T2B result; do not rebuild T2B from raw top-box rows unless no T2B table exists.'
    elif metric_type == 'multi_response':
        rules['total_answer'] = 'Use respondent-level incidence percentages; do not sum options because totals can exceed 100%.'
        rules['segment_answer'] = 'Compare each option independently; do not rank by summed multi-response percentages.'
    elif metric_type == 'mean_score':
        rules['total_answer'] = 'Use mean score with N; do not interpret it as a percentage distribution.'
        rules['significance_answer'] = 'Use mean-difference significance logic and compare means, not shares.'
    return rules


def _table_valid_comparisons(metric_type):
    comparisons = ['total_result', 'available_segment_breakdowns']
    if metric_type in ('distribution', 't2b', 'multi_response'):
        comparisons.append('percentage_point_differences')
    if metric_type == 'mean_score':
        comparisons.append('mean_score_differences')
    comparisons.append('available_wave_or_time_comparison')
    return comparisons


def _table_invalid_comparisons(metric_type, full_base_status):
    comparisons = ['missing_wave_as_zero', 'unavailable_breakdown_as_zero']
    if metric_type == 'multi_response':
        comparisons.append('summing_multi_response_options_to_100_percent')
    if metric_type == 'mean_score':
        comparisons.append('treating_mean_as_percentage')
    if full_base_status == 'partial_base':
        comparisons.append('comparing_partial_base_to_full_sample_without_routing_context')
    return comparisons


def _table_interpretation_limit_hints(metric_type, full_base_status):
    hints = []
    if full_base_status == 'partial_base':
        hints.append('Question has a reduced base; use routing_filter context before generalizing.')
    if metric_type == 'multi_response':
        hints.append('Multi-response percentages can exceed 100%; interpret options independently.')
    if metric_type == 't2b':
        hints.append('T2B is an aggregated positive/top category metric; avoid mixing with full distribution without saying so.')
    if metric_type == 'mean_score':
        hints.append('Mean scores need scale context; avoid percentage language.')
    return hints


def _add_ai_meta_warning(builder, **row):
    row.setdefault('audience', 'internal')
    row.setdefault('client_exclude', True)
    builder.add('WARNINGS', **row)


def _extract_q_code_from_title(title):
    """Return the lexical question code from a table title without interpreting it."""
    title = (title or '').strip()
    match = re.match(r'^([A-Za-z]+\d+[A-Za-z]*(?:\.\d+)?|[A-Za-z]+)', title)
    if not match:
        return ''
    return match.group(1).rstrip('.:').lower()


def _routing_group_key_from_title(title):
    title = (title or '').strip()
    match = re.match(r'^([A-Za-z]+\d+[A-Za-z]*)(?:[_\.][A-Za-z]?\d+)?', title)
    if match:
        return match.group(1).lower()
    return _extract_q_code_from_title(title)


def _lookup_value_label(vlabels, value):
    if not hasattr(vlabels, 'get'):
        return ''
    candidates = [value, str(value)]
    try:
        float_value = float(value)
        candidates.append(float_value)
        if float_value.is_integer():
            candidates.append(int(float_value))
    except (ValueError, TypeError):
        pass
    for candidate in candidates:
        try:
            label = vlabels.get(candidate, '')
        except TypeError:
            continue
        if label:
            return label
    return ''


def _labels_for_filter_values(filter_group, val_labels_dict):
    values = filter_group.get('vals', [])
    if filter_group.get('mode') == 'multi':
        group_vars = filter_group.get('vars', [])
        labels = []
        for value in values:
            label = ''
            for group_var in group_vars:
                label = _lookup_value_label(val_labels_dict.get(group_var, {}), value)
                if label:
                    break
            labels.append(label or str(value))
        return labels

    var_name = filter_group.get('var', '')
    vlabels = val_labels_dict.get(var_name, {})
    return [_lookup_value_label(vlabels, value) or str(value) for value in values]


def _label_or_var_name(labels_dict, var_name):
    label = labels_dict.get(var_name)
    if label is None:
        return str(var_name)
    try:
        if pd.isna(label):
            return str(var_name)
    except (TypeError, ValueError):
        pass
    label = str(label).strip()
    return label or str(var_name)


def _add_filter_meta_rows(builder, scope, output_id, filter_groups,
                          labels_dict, val_labels_dict):
    for condition_order, filter_group in enumerate(filter_groups, 1):
        mode = filter_group.get('mode', 'single')
        var_name = filter_group.get('var', '') if mode != 'multi' else ''
        group_vars = filter_group.get('vars', []) if mode == 'multi' else []
        builder.add(
            'FILTERS',
            scope=scope,
            output_id=output_id or '',
            condition_order=condition_order,
            logic_with_previous=filter_group.get('logic', 'AND'),
            mode=mode,
            var=var_name,
            vars_json=_json_cell(group_vars),
            group_label=filter_group.get('group_label', ''),
            var_label=_label_or_var_name(labels_dict, var_name) if var_name else '',
            value_codes_json=_json_cell(filter_group.get('vals', [])),
            value_labels_json=_json_cell(_labels_for_filter_values(filter_group, val_labels_dict)),
            description=build_filter_groups_description([filter_group], labels_dict, val_labels_dict),
        )


def _filter_description(filter_groups, labels_dict, val_labels_dict):
    return build_filter_groups_description(filter_groups, labels_dict, val_labels_dict) if filter_groups else ''


def _effective_filter_description(global_desc, output_desc):
    if global_desc and output_desc:
        return f'{global_desc} **I** {output_desc}'
    return global_desc or output_desc or ''


def _table_base_from_result(table):
    rows = table.get('rows', [])
    header = [str(h).strip().lower() for h in table.get('header', [])]
    for data_row in rows:
        if data_row and str(data_row[0]).strip().lower().startswith('total'):
            for candidate in ('n', 'frequency'):
                if candidate in header:
                    value = data_row[header.index(candidate)]
                    try:
                        return float(value), 'total_row'
                    except (TypeError, ValueError):
                        return '', 'total_row'
            if len(data_row) > 1:
                try:
                    return float(data_row[1]), 'total_row'
                except (TypeError, ValueError):
                    return '', 'total_row'

    if 'n' in header:
        n_idx = header.index('n')
        n_values = []
        for data_row in rows:
            if len(data_row) > n_idx:
                try:
                    n_values.append(float(data_row[n_idx]))
                except (TypeError, ValueError):
                    pass
        unique_values = sorted(set(n_values))
        if len(unique_values) == 1:
            return unique_values[0], 'numeric_n'
    return '', 'not_applicable'


def _table_base_frame(has_effective_filter, base_n, full_n, routing_known=False):
    if has_effective_filter:
        return 'output_filtered'
    if base_n == '':
        return 'not_applicable'
    try:
        if float(base_n) < float(full_n):
            if routing_known:
                return 'output_filtered'
            return 'partial_unknown'
    except (TypeError, ValueError):
        pass
    return 'full'


def _output_routing_filter_status(routing_meta, base_frame):
    if _routing_known(routing_meta):
        return 'known'
    if base_frame == 'partial_unknown':
        return 'unknown_partial_base'
    return 'not_filtered'


def _estimate_definition_base(df, table_type, variables_resolved):
    if not variables_resolved:
        return '', 'unknown', 'not_filtered'
    try:
        if table_type in ('s', 'f'):
            valid_n = int(df[variables_resolved[0]].notna().sum())
        elif table_type in ('k', 'd'):
            valid_n = int(df[variables_resolved].notna().any(axis=1).sum())
        elif table_type in ('n', 'm'):
            return '', 'not_applicable', 'not_filtered'
        else:
            return '', 'unknown', 'not_filtered'
    except Exception:
        return '', 'unknown', 'not_filtered'

    if valid_n < len(df):
        return valid_n, 'partial_base', 'unknown_partial_base'
    return valid_n, 'full_sample', 'not_filtered'


def _routing_meta_for_table(ai_meta_settings, table_idx):
    if not ai_meta_settings:
        return {}
    return ai_meta_settings.get('routing', {}).get(str(table_idx), {})


def _routing_known(routing_meta):
    return bool((routing_meta.get('description') or '').strip() or
                (routing_meta.get('expression') or '').strip())


def _partial_base_candidates(titles, variables, df, meta, start_num):
    from collections import OrderedDict

    labels_dict = getattr(meta, 'column_names_to_labels', {}) or {}
    col_lc = {col.lower(): col for col in df.columns}
    groups = OrderedDict()
    for table_idx, (title_line, var_line) in enumerate(zip(titles, variables)):
        table_type = get_table_type(title_line)
        table_title = get_table_title(title_line)
        raw_vars = _extract_vars_from_line(var_line)
        variables_resolved = [col_lc[var.lower()] for var in raw_vars if var.lower() in col_lc]
        full_base_n, full_base_status, _routing_status = _estimate_definition_base(
            df, table_type, variables_resolved)
        if full_base_status != 'partial_base':
            continue
        labels = [_label_or_var_name(labels_dict, var) for var in variables_resolved]
        q_code = _routing_group_key_from_title(table_title)
        group = groups.setdefault(q_code or f'table_{table_idx}', {
            'table_idx': table_idx,
            'table_indices': [],
            'table_numbers': [],
            'q_code': q_code,
            'title': table_title,
            'variables': [],
            'variable_labels': [],
        })
        group['table_indices'].append(table_idx)
        group['table_numbers'].append(f'{table_idx + start_num}.1')
        for var_name, var_label in zip(variables_resolved, labels):
            if var_name not in group['variables']:
                group['variables'].append(var_name)
                group['variable_labels'].append(var_label)

    candidates = []
    for group in groups.values():
        group_vars = group['variables']
        if group_vars:
            base_n = int(df[group_vars].notna().any(axis=1).sum())
        else:
            base_n = ''
        table_numbers = group['table_numbers']
        group['table_number'] = table_numbers[0] if len(table_numbers) == 1 else f'{table_numbers[0]} - {table_numbers[-1]}'
        group['base_n'] = base_n
        candidates.append(group)
    return candidates


def _add_table_definition_meta(builder, titles, variables, df, meta, start_num,
                               ai_meta_settings=None):
    labels_dict = getattr(meta, 'column_names_to_labels', {}) or {}
    val_labels_dict = getattr(meta, 'variable_value_labels', {}) or {}
    col_lc = {col.lower(): col for col in df.columns}
    for table_idx, (title_line, var_line) in enumerate(zip(titles, variables)):
        table_type = get_table_type(title_line)
        table_title = get_table_title(title_line)
        raw_vars = _extract_vars_from_line(var_line)
        variables_resolved = [col_lc[var.lower()] for var in raw_vars if var.lower() in col_lc]
        variable_labels = [_label_or_var_name(labels_dict, var) for var in variables_resolved]
        full_base_n, full_base_status, routing_status = _estimate_definition_base(
            df, table_type, variables_resolved)
        title_upper = table_title.upper()
        routing_meta = _routing_meta_for_table(ai_meta_settings, table_idx)
        routing_description = _normalize_user_text(routing_meta.get('description')).strip()
        routing_expression = _normalize_user_text(routing_meta.get('expression')).strip()
        if _routing_known(routing_meta):
            routing_status = 'known'
        metric_type = _table_metric_type(table_type, table_title)
        result_level = _table_result_level(full_base_status)

        builder.add(
            'TABLES',
            table_idx=table_idx,
            table_number=f'{table_idx + start_num}.1',
            q_code=_extract_q_code_from_title(table_title),
            title=table_title,
            input_title_line=title_line,
            input_var_line=var_line,
            table_type_code=table_type,
            table_type_label=_TABLE_TYPE_LABELS.get(table_type, table_type),
            variables_json=_json_cell(variables_resolved),
            variable_labels_json=_json_cell(variable_labels),
            metric_type=metric_type,
            result_level=result_level,
            answer_scale_json=_json_cell(_answer_scale_meta(metric_type, variables_resolved, val_labels_dict)),
            evidence_needed_rules_json=_json_cell(_table_evidence_needed_rules(metric_type)),
            valid_comparisons_json=_json_cell(_table_valid_comparisons(metric_type)),
            invalid_comparisons_json=_json_cell(_table_invalid_comparisons(metric_type, full_base_status)),
            interpretation_limit_hints_json=_json_cell(_table_interpretation_limit_hints(metric_type, full_base_status)),
            is_multi_response=table_type in ('k', 'd'),
            is_mean=table_type in ('n', 'm') or '- MEAN' in title_upper,
            is_t2b='T2B' in title_upper,
            full_base_n=full_base_n,
            full_base_status=full_base_status,
            routing_filter_description=routing_description,
            routing_filter_expression=routing_expression,
            routing_filter_source='user' if _routing_known(routing_meta) else '',
            routing_filter_status=routing_status,
        )
        if _routing_known(routing_meta):
            builder.add(
                'ROUTING',
                table_idx=table_idx,
                table_number=f'{table_idx + start_num}.1',
                q_code=_extract_q_code_from_title(table_title),
                title=table_title,
                base_n=full_base_n,
                description=routing_description,
                expression=routing_expression,
                source='user',
            )


def _add_banner_meta_rows(builder, output_id, sheet_name, banner_vars, work_df,
                          meta, weight_col):
    val_labels_dict = getattr(meta, 'variable_value_labels', {}) or {}
    for banner_order, banner_var in enumerate(banner_vars, 1):
        if banner_var not in work_df.columns:
            _add_ai_meta_warning(
                builder,
                level='warning', code='BANNER_VAR_MISSING',
                table_idx='', output_id=output_id,
                message=f'Banner variable does not exist: {banner_var}',
            )
            continue
        banner_label = get_var_label(banner_var, meta)
        if banner_label == banner_var:
            _add_ai_meta_warning(
                builder,
                level='warning', code='BANNER_LABEL_MISSING',
                table_idx='', output_id=output_id,
                message=f'Banner variable has no SAV label; falling back to variable name: {banner_var}',
            )
        value_labels = val_labels_dict.get(banner_var, {})
        observed_values = sorted(
            work_df[banner_var].dropna().unique(),
            key=lambda value: (0, float(value)) if isinstance(value, (int, float, np.integer, np.floating)) else (1, str(value)),
        )
        for segment_order, segment_code in enumerate(observed_values, 1):
            segment_label = _lookup_value_label(value_labels, segment_code) or str(segment_code)
            segment_mask = work_df[banner_var] == segment_code
            if weight_col and weight_col in work_df.columns:
                segment_n = float(work_df.loc[segment_mask, weight_col].sum())
            else:
                segment_n = int(segment_mask.sum())
            builder.add(
                'BANNERS',
                output_id=output_id,
                sheet_name=sheet_name,
                banner_order=banner_order,
                banner_var=banner_var,
                banner_label=banner_label,
                axis_id=banner_var,
                segment_order=segment_order,
                segment_code=segment_code,
                segment_label=segment_label,
                segment_n=segment_n,
                weight_col=weight_col or '',
            )


# ═══════════════════════════════════════════════════════════════════
#  THE WRITER — the hooks app.py and headless.py call, in loop order
# ═══════════════════════════════════════════════════════════════════

AI_META_SHEET_NAME = '_AI_META'
META_SCHEMA_VERSION = '1.2'
APP_NAME = 'Hendalice'
APP_VERSION = '2.5'


def infer_project_code(sav_name):
    """'Data_online_26_04_049_divir2_v5.sav' -> ('online_26_04_049', 'filename'); ('', 'unknown') when absent."""
    stem = os.path.splitext(sav_name or '')[0]
    match = re.search(r'([a-z]+_\d{2}_\d{2}_\d{3})', stem, re.IGNORECASE)
    if match:
        return match.group(1).lower(), 'filename'
    return '', 'unknown'


class AiMetaWriter:
    """Feeds the _AI_META rows from the generate loop; one instance per workbook.

    Call order (both entry points, identical):
        w = AiMetaWriter.start(df, meta, titles, variables, plan=..., ...)   once, before the outputs
        w.begin_output(output_id, out_def, n_after_global, n_after_effective) per output, after its filters
        w.add_total_sheet(output_id, sheet_name, out_def, tables)            per copied total sheet
        w.add_between_options(output_id, source_sheet, between_sheet, blocks) when the _btw sheet was written
        w.add_krizanje_output(...)                                           per krizanje output, sheets created
        w.add_banner_table(output_id, sheet_name, role, entry, start, end)  per entry x (cross_base|significance|sig_total)
        w.finish(wb)                                                         last, after the TOC

    `plan` is the po.json `global.ai_meta` dict ({'study': {...}, 'routing': {idx: {...}}}) —
    app.py builds it from the form (_collect_ai_meta_plan), headless.py reads it from the file.
    """

    def __init__(self, builder, df, meta, settings, use_weight, weight_col, start_num):
        self.builder = builder
        self.df = df
        self.n_unfiltered = len(df)
        self.meta = meta
        self.settings = settings
        self.use_weight = use_weight
        self.weight_col = weight_col
        self.start_num = start_num
        self.labels_dict = getattr(meta, 'column_names_to_labels', {}) or {}
        self.val_labels_dict = getattr(meta, 'variable_value_labels', {}) or {}
        self.global_filter_groups = []
        self._outputs = {}

    @classmethod
    def start(cls, df, meta, titles, variables, *, plan=None, use_weight=False, weight_col=None,
              start_num=1, table_design='hendal', sav_name='', input_name='',
              global_filter_groups=None, generated_at=None):
        plan = plan or {}
        settings = {
            'study': dict(plan.get('study', {}) or {}),
            'routing': dict(plan.get('routing', {}) or {}),
        }
        builder = AiMetaBuilder()
        writer = cls(builder, df, meta, settings, use_weight, weight_col, start_num)
        writer.global_filter_groups = list(global_filter_groups or [])

        b = builder
        b.add('META', key='meta_schema_version', value=META_SCHEMA_VERSION, source='app', notes='')
        b.add('META', key='generated_at',
              value=generated_at or datetime.now(timezone.utc).isoformat(), source='app', notes='')
        b.add('META', key='app_name', value=APP_NAME, source='app', notes='')
        b.add('META', key='app_version', value=APP_VERSION, source='app', notes='')
        b.add('META', key='source_sav_name', value=sav_name or '', source='sav', notes='')
        b.add('META', key='source_input_name', value=input_name or '', source='input', notes='')
        b.add('META', key='row_count_unfiltered', value=len(df), source='dataframe', notes='')
        b.add('META', key='weight_enabled', value=bool(use_weight), source='app', notes='')
        b.add('META', key='weight_col', value=weight_col or '', source='app', notes='')
        b.add('META', key='start_num', value=start_num, source='app', notes='')
        b.add('META', key='table_design', value=table_design, source='app', notes='')
        project_code, project_code_source = infer_project_code(sav_name)
        b.add('META', key='project_code_inferred', value=project_code, source=project_code_source, notes='')
        _add_study_meta_rows(b, settings['study'], df, use_weight, weight_col)

        _add_table_definition_meta(b, titles, variables, df, meta, start_num, ai_meta_settings=settings)
        _add_filter_meta_rows(b, 'global', '', writer.global_filter_groups,
                              writer.labels_dict, writer.val_labels_dict)
        return writer

    # ── per output ──

    def begin_output(self, output_id, out_def, n_after_global_filter, n_after_effective_filter):
        output_filter_groups = out_def.get('filter_groups', []) or []
        global_desc = _filter_description(self.global_filter_groups, self.labels_dict, self.val_labels_dict)
        output_desc = _filter_description(output_filter_groups, self.labels_dict, self.val_labels_dict)
        self._outputs[output_id] = {
            'global_desc': global_desc,
            'output_desc': output_desc,
            'effective_desc': _effective_filter_description(global_desc, output_desc),
            'n_after_global': n_after_global_filter,
            'n_after_effective': n_after_effective_filter,
            'filter_groups': output_filter_groups,
            'has_filter': bool(self.global_filter_groups or output_filter_groups),
        }
        _add_filter_meta_rows(self.builder, 'output', output_id, output_filter_groups,
                              self.labels_dict, self.val_labels_dict)

    def _effective_groups_json(self, output_id):
        return _json_cell({
            'global': self.global_filter_groups,
            'output': self._outputs[output_id]['filter_groups'],
        })

    def _add_outputs_row(self, output_id, out_def, output_type, actual_base, actual_sig, actual_sig_total,
                         banner_vars, banner_labels_json, show_sig, show_sig_total):
        o = self._outputs[output_id]
        table_indices = out_def.get('table_indices', [])
        self.builder.add(
            'OUTPUTS',
            output_id=output_id,
            sheet_base_name=out_def['sheet_name'],
            output_type=output_type,
            actual_sheet_base=actual_base,
            actual_sheet_sig=actual_sig,
            actual_sheet_sig_total=actual_sig_total,
            selected_table_indices_json=_json_cell(table_indices),
            selected_table_count=len(table_indices),
            banner_vars_json=_json_cell(banner_vars),
            banner_labels_json=banner_labels_json,
            show_sig=show_sig,
            show_sig_total=show_sig_total,
            global_filter_description=o['global_desc'],
            output_filter_description=o['output_desc'],
            effective_filter_description=o['effective_desc'],
            filter_groups_json=_json_cell(out_def.get('filter_groups', [])),
            global_filter_groups_json=_json_cell(self.global_filter_groups),
            n_unfiltered=self.n_unfiltered,
            n_after_global_filter=o['n_after_global'],
            n_after_effective_filter=o['n_after_effective'],
            weight_col=self.weight_col or '',
        )

    # ── total output ──

    def add_total_sheet(self, output_id, sheet_name, out_def, tables):
        b = self.builder
        o = self._outputs[output_id]
        b.add(
            'SHEETS', sheet_name=sheet_name, role='total',
            output_id=output_id, base_sheet_name='',
            table_block_count=len(tables), hidden=False,
        )
        self._add_outputs_row(output_id, out_def, 'total', sheet_name, '', '',
                              [], _json_cell([]), False, False)

        meta_row = 1
        for tbl in tables:
            n_data = len(tbl['rows'])
            has_caption = bool(tbl.get('caption', ''))
            start_row = meta_row
            end_row = start_row + 1 + 1 + n_data + (1 if has_caption else 0) - 1
            base_n, base_source = _table_base_from_result(tbl)
            table_idx = tbl.get('_idx', '')
            routing_meta = _routing_meta_for_table(self.settings, table_idx) if table_idx != '' else {}
            base_frame = _table_base_frame(
                o['has_filter'], base_n, self.n_unfiltered, routing_known=_routing_known(routing_meta),
            )
            b.add(
                'OUTPUT_TABLES',
                output_id=output_id,
                sheet_name=sheet_name,
                sheet_role='total',
                table_idx=table_idx,
                table_number=f'{int(table_idx) + self.start_num}.1' if table_idx != '' else '',
                title_rendered=tbl.get('title', ''),
                start_row=start_row,
                end_row=end_row,
                base_n=base_n,
                base_source=base_source,
                effective_filter_description=o['effective_desc'],
                effective_filter_groups_json=self._effective_groups_json(output_id),
                routing_filter_description=routing_meta.get('description', ''),
                routing_filter_expression=routing_meta.get('expression', ''),
                routing_filter_status=_output_routing_filter_status(routing_meta, base_frame),
                base_frame=base_frame,
                banner_vars_json=_json_cell([]),
            )
            if base_frame == 'partial_unknown':
                _add_ai_meta_warning(
                    b,
                    level='warning',
                    code='PARTIAL_BASE_UNKNOWN_FILTER',
                    table_idx=table_idx, output_id=output_id,
                    message='Table base is lower than full sample and no explicit output/global filter is active.',
                )
            meta_row += 1 + 1 + n_data + (1 if has_caption else 0) + 1

    def add_between_options(self, output_id, source_sheet, between_sheet, between_blocks):
        b = self.builder
        b.add(
            'SHEETS', sheet_name=between_sheet, role='between_options',
            output_id=output_id, base_sheet_name=source_sheet,
            table_block_count=len(between_blocks), hidden=False,
        )
        for blk in between_blocks:
            b_idx = blk['table_idx']
            b_metric = _table_metric_type(blk['table_type'], blk['title'])
            b_tnum = f'{int(b_idx) + self.start_num}.1'
            for pair in blk['results']:
                b.add(
                    'BETWEEN_OPTIONS',
                    output_id=output_id,
                    source_sheet=source_sheet,
                    between_sheet=between_sheet,
                    table_idx=b_idx,
                    table_number=b_tnum,
                    q_code=blk['q_code'],
                    title=blk['title'],
                    table_type_code=blk['table_type'],
                    metric_type=b_metric,
                    scope='total',
                    axis='Total',
                    base_n=pair['n'],
                    option_a_label=pair['option_a_label'],
                    option_a_value=pair['option_a_value'],
                    option_b_label=pair['option_b_label'],
                    option_b_value=pair['option_b_value'],
                    significant=pair['significant'],
                    direction=pair['direction'],
                    test=pair['test'],
                    confidence=pair['confidence'],
                    b_a_not_b=pair.get('b_a_not_b', ''),
                    c_b_not_a=pair.get('c_b_not_a', ''),
                    note=pair.get('note', ''),
                )

    # ── krizanje output ──

    def add_krizanje_output(self, output_id, out_def, ws_name, sig_name, sig_total_name, n_entries,
                            banner_vars, banner_labels, show_sig, show_sig_total, work_df):
        b = self.builder
        b.add(
            'SHEETS', sheet_name=ws_name, role='cross_base',
            output_id=output_id, base_sheet_name='',
            table_block_count=n_entries, hidden=False,
        )
        if sig_name:
            b.add(
                'SHEETS', sheet_name=sig_name, role='significance',
                output_id=output_id, base_sheet_name=ws_name,
                table_block_count=n_entries, hidden=False,
            )
        if sig_total_name:
            b.add(
                'SHEETS', sheet_name=sig_total_name, role='sig_total',
                output_id=output_id, base_sheet_name=ws_name,
                table_block_count=n_entries, hidden=False,
            )
        self._add_outputs_row(output_id, out_def, 'krizanje', ws_name, sig_name or '', sig_total_name or '',
                              banner_vars, _json_cell(banner_labels or []),
                              bool(show_sig), bool(show_sig_total))
        _add_banner_meta_rows(b, output_id, ws_name, banner_vars, work_df, self.meta, self.weight_col)

    def add_banner_table(self, output_id, sheet_name, sheet_role, entry, start_row, end_row):
        """One OUTPUT_TABLES row per written banner block; the partial-base warning only on the base sheet."""
        b = self.builder
        o = self._outputs[output_id]
        ti = entry['table_idx']
        banner = entry['banner']
        routing_meta = _routing_meta_for_table(self.settings, ti)
        base_frame = _table_base_frame(
            o['has_filter'], banner.get('total_n', ''), self.n_unfiltered,
            routing_known=_routing_known(routing_meta),
        )
        b.add(
            'OUTPUT_TABLES',
            output_id=output_id,
            sheet_name=sheet_name,
            sheet_role=sheet_role,
            table_idx=ti,
            table_number=f"{entry['table_num']}.1",
            title_rendered=entry['title'],
            start_row=start_row,
            end_row=end_row,
            base_n=banner.get('total_n', ''),
            base_source='n_row',
            effective_filter_description=o['effective_desc'],
            effective_filter_groups_json=self._effective_groups_json(output_id),
            routing_filter_description=routing_meta.get('description', ''),
            routing_filter_expression=routing_meta.get('expression', ''),
            routing_filter_status=_output_routing_filter_status(routing_meta, base_frame),
            base_frame=base_frame,
            banner_vars_json=_json_cell(entry['banner_vars']),
        )
        if sheet_role == 'cross_base' and base_frame == 'partial_unknown':
            _add_ai_meta_warning(
                b,
                level='warning',
                code='PARTIAL_BASE_UNKNOWN_FILTER',
                table_idx=ti, output_id=output_id,
                message='Crosstab base is lower than full sample and no explicit output/global filter is active.',
            )

    # ── end ──

    def finish(self, wb):
        self.builder.add(
            'SHEETS', sheet_name=AI_META_SHEET_NAME, role='ai_meta',
            output_id='', base_sheet_name='', table_block_count=0,
            hidden=False,
        )
        return self.builder.write_to_workbook(wb, sheet_name=AI_META_SHEET_NAME, hidden=False)
